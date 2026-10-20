"""Update only Chapter 5 text/tables from a sealed evidence run.

Usage: python revise-paper85-section5.py SOURCE.docx RUN_DIR OUTPUT.docx
The adjacent claims JSON records inputs, hashes, selectors and displayed cells.
"""
import copy
import hashlib
import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.shared import Pt

SOURCE_SHA = '1515cd7f87995759b4a2f047aa314e23aae9cd378e25ae80b5790a446fca252e'
source, run_dir, output = (Path(p).resolve() for p in sys.argv[1:])
root = Path(__file__).resolve().parents[1]
hash_file = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
load = lambda p: json.loads(p.read_text())
assert hash_file(source) == SOURCE_SHA, 'Unexpected source manuscript; review before applying indexed edits'
manifest = load(run_dir / 'manifest.json')
assert manifest['status'] == 'complete'
for path, digest in manifest['files'].items():
    assert hash_file(run_dir / path) == digest, path

A = load(run_dir / 'A-summary.json')
B = load(run_dir / 'B-scale-summary.json')
C = load(run_dir / 'C-case-summary.json')
D = load(run_dir / 'D-posthoc-revoked-transfer.json')
E = load(run_dir / 'E-next-nav-counterexample.json')
host = load(run_dir / 'run.json')
old_dir = root / 'docs/evidence/paper85/v2.1'
m0 = load(old_dir / 'm0-run-summary.local.json')
abl = load(old_dir / 'ablation-run-summary.local.json')
assert m0['counts']['m0Equivalent'] == 42 and abl['counts']['total'] == 210
assert C['answers']['Q1']['matchingCasePolicy'] == 4 and C['answers']['Q1']['total'] == 5
assert E['pricing'][1]['reason'] == 'skipped-next-NAV'

doc = Document(source)
blocks = [e for e in doc._element.body if e.tag == qn('w:tbl')
          or (e.tag == qn('w:p') and ''.join(e.xpath('.//w:t/text()')).strip())]
assert len(blocks) == 154
original_text = [''.join(e.itertext()) for e in blocks[:98] + blocks[152:]]
claims = {'sourceSha256': SOURCE_SHA, 'run': str(run_dir.relative_to(root)),
          'manifestSha256': hash_file(run_dir / 'manifest.json'), 'executionCommit': host['commit'],
          'generatorSha256': hash_file(Path(__file__)),
          'sourceResolution': 'Run-relative paths first; spec-v3 and scripts are relative to experiments/paper85-evaluation; contracts and experiments paths are repository-relative. Wildcards select all matching archived evidence.',
          'blocks': {}, 'tables': {}}


def replace_paragraph(p, text):
    # Retain paragraph layout and the first run's typography.
    rpr = copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    run = p.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)


def paragraph(index, text, sources=()):
    replace_paragraph(Paragraph(blocks[index], doc), text)
    claims['blocks'][str(index)] = {'sources': list(sources), 'text': text}


def cell(table_index, row, col, text):
    t = Table(blocks[table_index], doc)
    c = t.cell(row, col)
    replace_paragraph(c.paragraphs[0], str(text))
    for p in list(c.paragraphs)[1:]:
        p._element.getparent().remove(p._element)


def table_record(index, title, sources):
    t = Table(blocks[index], doc)
    claims['tables'][title] = {'sources': sources, 'displayedCells': [[c.text for c in r.cells] for r in t.rows]}


def n(value, digits=0):
    return f'{value:,.{digits}f}'


def ms(metric, digits=0, scale=1):
    return f"{n(metric['mean'] / scale, digits)} ± {n(metric['sd'] / scale, digits)}"


paragraph(99, '本节围绕五个评价问题（evaluation question，EQ）组织原型评价（表3）。EQ1采用功能与机制验证，EQ2结合消融与两类自建比较基线，EQ3采用闭合输入回放及检查点比对，EQ4测量执行成本与顺序负载下的规模增长，EQ5采用公开规则驱动的案例验证。原始场景与v3评价方案保持冻结；审查后修正的NAV审计判据、拒绝分类、证据归档、构建校验和负载时钟另列为v3.1补充方案，明确属于事后修订。本文区分确定性实现检查、描述性性能测量和案例相关性证据，不将其等同于生产验证。', ['spec-v3/evaluation-plan.v3.md', 'spec-v3/review-amendment.v3.1.md'])
cell(101, 4, 2, '交易回执gas与比较基线；100、1,000、10,000事件顺序规模增长实验')
cell(101, 5, 1, '所选公开规则下的设计目标与生命周期映射是否相关，暴露哪些完整性缺口')
paragraph(103, '原型由三份Solidity合约构成。FundToken维护份额、资格、申购赎回请求、待赎回锁定份额及运营状态；NAVRegistry维护追加式NAV序列；RiskRegistry提供赎回控制接口，本评价仅配置该依赖，不提交风险指标或触发风险控制。实现采用OpenZeppelin的ERC-20、AccessControl与Pausable组件。生产源码与chapter3-artifact-v1.4.0逐字节一致。本次运行先核对冻结源码与原预注册，再重建并以归档的字面源码和设置独立编译，比较创建及运行时字节码；部署后再次核对代码（构造器不可变值位置单独保留）。表4报告本次实际环境和构建结果。', ['run.json#/files', 'build/*.input.json.gz', 'build/*.output.json.gz', 'build/*.artifact.json', 'chains/*.jsonl.gz'])
cell(105, 1, 1, 'Solidity；solc 0.8.30+commit.73712a01；Osaka；optimizer=false、runs=200；IPFS元数据；完整设置、源码及编译输出随结果归档')
cell(105, 4, 1, f"Node.js {host['node'].lstrip('v')}，viem {A['provenance']['tools']['viem']}；纯回放只读取创世配置和日志，合约视图独立保存用于结果比对；规模检查点视图在负载生成时采集，回放不读取视图")
hw = host['hardware']
cell(105, 5, 1, f"{hw['cpu']}，{hw['logicalCpus']}个逻辑CPU，{hw['memoryBytes'] / 1024**3:.0f} GiB内存，{hw['platform']} / Darwin {hw['release']}；套件顺序执行，未保证主机独占；时延为本机描述性观测")
cell(105, 6, 1, '标准轨迹5个活动账户；规模27个活动账户（24名投资者，链预资助28个账户）；案例7个活动账户（预资助8个）。FundToken含管理员、申购、赎回、暂停角色；NAVRegistry含估值管理角色')
artifacts = {name: load(run_dir / f'build/{name}.artifact.json') for name in ['FundToken', 'NAVRegistry', 'RiskRegistry', 'B1TokenOnly', 'B2PermissionedAllowlist']}
sizes = {name: (len(a['deployedBytecode']['object']) - 2) // 2 for name, a in artifacts.items()}
cell(105, 7, 1, '；'.join(f"{name} {n(sizes[name])} B" for name in sizes))
deployment = dict(zip(['NAVRegistry', 'FundToken', 'RiskRegistry'], A['productionGas']['deploymentGas']))
deployment.update({kind: A['comparatorGas'][kind]['deploymentGas'][0] for kind in ['B1', 'B2']})
cell(105, 8, 1, '；'.join(f'{name} {n(value)}' for name, value in deployment.items()) + '；B2另有3笔角色授予。采用本次重建的实测值，不以源码不变推定部署gas等于历史值')
cell(105, 9, 1, f"读取成功交易回执gasUsed；生产、B1、B2各{A['productionGas']['repetitions']}次全新标准轨迹；规模为3个种子×3个目标；获取日志5次，纯回放1次预热后10次计时。拒绝项为eth_call模拟，不是已上链的失败交易")
paragraph(106, '原型对第3节模型采用窄化实现，结论限于表5覆盖的语义。资格采用布尔白名单，未实现冻结、复核和受控退出；请求没有显式NAV版本参数，受理与结算使用当时最新登记版本，适用性不在链上判断；转让只检查接收方资格。请求记录与待赎回锁定份额是两阶段实现的辅助状态，属于抽象模型到实现的细化；本评价检查具体实现投影，不证明原型完整实现抽象模型。', ['contracts/src/FundToken.sol', 'contracts/src/NAVRegistry.sol'])
paragraph(109, f"42场景与单谓词消融沿用v2.1冻结证据（运行实现d27d990）；原v3预注册提交4ce5d24及其结果完整保留。本次修订运行{run_dir.name}，执行代码提交{host['commit'][:7]}，新增结果均取自该独立目录。运行清单绑定源码、构建、调用、回执、原始日志、检查点视图和时延样本；离线验证不连接RPC。逐表来源、JSON路径与显示值见随稿证据映射paper85-section5-claims.json。哈希用于内部溯源和完整性校验，不提供独立机构背书。", ['run.json', 'manifest.json', 'verification.json', '../../v2.1/m0-run-summary.local.json', '../../v2.1/ablation-run-summary.local.json'])
paragraph(114, '18个有效场景被接受，24个无效场景按预期错误码拒绝且状态投影不变，STCR = 42/42 = 1.00；全约束实验变体与生产合约在42个场景上差分等价（表6）。冻结证据还包含本地与CI的语义一致性检查。确定性实现通过确定性测试是预期结果，STCR本身不构成强经验证据。结果支持原型在所选条件下符合具体转换语义，其覆盖范围受场景集和表5中的模型差距限制。', ['../../v2.1/m0-run-summary.local.json#/counts', '../../v2.1/m0-run-summary.ci.json#/counts'])
paragraph(116, 'EQ2保留五类单谓词消融，每次只移除一个谓词检查，其余逻辑不变，运行全部42个场景。移除guard后放行相应操作具有同义反复，因此消融用于检查剩余约束能否在这些场景中阻止同类无效转换，不证明一般必要性。表7中24个目标场景有22个被接受；I10在M3下仍因无NAV记录回退，M5中另有一个场景触发算术下溢。全部210项分类与预期一致，24个目标场景的62项字段断言通过。两类独立实现的比较基线进一步检验原型相对常规代币账本所增加的约束范围。', ['../../v2.1/ablation-run-summary.local.json#/counts'])
cell(121, 0, 2, '原型（表5范围）')
paragraph(122, '注：每格依次为“拒绝 / 接受且改变余额 / 接受但不改变余额 / 无法在链上表达”的场景数。B1/B2的拒绝由eth_call预检确认EVM回退，接受项读取已上链交易的回执及前后投影；编码、网络、超时及其他RPC故障中止运行，不计为拦截。零数量铸造或销毁可成功写入零值Transfer记录。完整机制结果沿用表6。', ['A-scenario-rows.json', 'chains/*.jsonl.gz'])
paragraph(125, 'EQ3检查事件能否独立于合约视图恢复创世配置预先声明的账户与角色范围内的实现投影。标准轨迹包含14个操作、22个生命周期事件，覆盖资格、NAV、请求、份额变动、延迟、暂停恢复与角色变化。纯回放仅输入创世配置和有序原始日志，合约历史视图另行保存作为比对结果。本次ARC = 22/22、NRI = 2/2、PCR = 14/14；NRI仅表示引用绑定执行时最新登记版本且随后可恢复，不表示该版本在业务或监管上适用。基线恢复结果见表9；原始日志、ABI、创世配置和各检查点均归档。', ['canonical/18660.json.gz', 'A-summary.json#/recoverability', 'verification.json'])
paragraph(126, Paragraph(blocks[126], doc).text + ' B2的执行角色项只比较候选发行人权限与原型申购执行权的对应，不代表完整机构权限模型。', ['A-summary.json#/recoverability'])
nri = sum(x['total'] for a in B['aggregate'] for x in a['nri'])
nri10 = sum(x['total'] for x in B['aggregate'][-1]['nri'])
pcr = sum(x['checkpoints'] for a in B['aggregate'] for x in a['pcr'])
paragraph(129, f"全部9次规模运行的ARC均为1.00，NRI合计{n(nri)}/{n(nri)}（10,000事件目标下{n(nri10)}/{n(nri10)}），PCR合计{pcr}/{pcr}。案例的ARC = {C['arc']['replayed']}/{C['arc']['accepted']}、NRI = {C['nri']['valid']}/{C['nri']['total']}、PCR = {C['answers']['Q6']['matched']}/{C['answers']['Q6']['checkpoints']}。这些结果依赖可信创世配置、完整有序日志及无链重组。规模检查点的视图与回放输入分开保存，回放不消费视图。模拟拒绝不产生生命周期事件，需要另存调用及EVM错误；本实验不将其称为失败交易。受理与结算执行者仍需交易信封补充，案例政策也属于审计的外部假设。", ['B-scale-summary.json#/aggregate', 'C-case-summary.json', 'scale/*/replay.json.gz'])
paragraph(131, f"EQ4先比较标准轨迹上的交易gas（表10）。生产合约与两类比较基线各执行{A['productionGas']['repetitions']}次全新部署；生产各操作gas在重复运行中一致。所有值从本次保存的交易回执归集，完整机制的申购和赎回包含请求与生效两笔交易，基线只执行一次铸造或销毁。", ['A-summary.json#/productionGas', 'A-summary.json#/comparatorGas', 'chains/*.jsonl.gz'])
op = {key: value['meanGas'] for key, value in A['productionGas']['operations'].items()}
categories = [('资格更新', ['O01', 'O02'], 'qualification', False), ('NAV提交', ['O03', 'O07'], 'navSubmission', False),
              ('申购', ['O04', 'O05'], 'subscription', True), ('赎回', ['O08', 'O10'], 'redemption', True),
              ('受限转让', ['O06'], 'transfer', False), ('暂停/恢复', ['O11', 'O12'], 'pauseResume', False),
              ('角色更新', ['O13', 'O14'], 'roleUpdate', False), ('结算延迟标记', ['O09'], 'settlementDelay', False)]
for row, (_, ids, key, summed) in enumerate(categories, 1):
    value = sum(op[i] for i in ids) / (1 if summed else len(ids))
    text = n(value, 1 if value % 1 else 0)
    if summed:
        text += f"（请求{n(op[ids[0]])}＋{'接受' if key == 'subscription' else '结算'}{n(op[ids[1]])}）"
    cell(133, row, 1, text)
    for col, kind in [(2, 'B2'), (3, 'B1')]:
        baseline = A['comparatorGas'][kind]['operations'].get(key)
        cell(133, row, col, n(baseline['meanGas']) if baseline else '—')
    b2 = A['comparatorGas']['B2']['operations'].get(key)
    cell(133, row, 4, n(value / b2['meanGas'], 2) if b2 else '—')
paragraph(136, '第二部分为单节点顺序规模增长实验。24名投资者的混合负载由固定种子生成，只在当前可执行操作中按原方案权重抽样。各权重为：申购请求14、接受14、转让18、赎回请求12、结算12、延迟2、NAV提交10、资格更新8、暂停恢复对4、角色变更对2。一次操作一个区块，时间推进60秒；NAV估值时间与其提交区块时间相同并逐项断言。事件数达到100、1,000或10,000时在操作边界停止，因此允许超过目标。各目标使用3个种子，表11报告种子间均值与样本标准差。该设计不测并发饱和、节点竞争或故障恢复。', ['B-scale-summary.json#/provenance', 'scale/*/replay.json.gz'])
paragraph(137, '表11 顺序规模增长结果（3个种子，均值 ± 样本标准差）')
scale_rows = [('transactions', 0, 1), ('acceptedEvents', 0, 1), ('gasPerTransaction', 0, 1), ('gasPerEvent', 0, 1),
              ('storageSlots', 0, 1), ('storageSlotsPerTransaction', 2, 1), ('rawPayloadBytes', 1, 1024),
              ('rawPayloadBytesPerEvent', 1, 1), ('normalizedJsonlBytes', 1, 1024), ('getLogsMs', 2, 1),
              ('replayMs', 2, 1), ('replayUsPerEvent', 2, 1)]
for row, (key, digits, scale) in enumerate(scale_rows, 1):
    for col, aggregate in enumerate(B['aggregate'], 1):
        cell(138, row, col, ms(aggregate[key], digits, scale))
for row in [7, 9]:
    c = Table(blocks[138], doc).cell(row, 0)
    replace_paragraph(c.paragraphs[0], c.text.replace('KB', 'KiB'))
paragraph(139, '注：样本标准差采用n−1分母，时延表项先求每次运行的样本均值，再对3个种子求均值与标准差。原始载荷为每个topic 32字节加data字节；KiB = 1,024字节。事件索引为归档的未压缩JSONL字符串。获取日志按2,000区块分页、重复5次；纯回放预热1次、计时10次，不含获取、解码、最终引用检查或证据写入，各次时延样本全部保存。存储指标为三份合约运行末的非零槽总数，包含初始化状态；槽数/交易数不是净新增槽数/操作，不代表磁盘字节数。', ['B-scale-summary.json#/aggregate', 'scale/*/replay.json.gz#/timingSamplesMs', 'chains/*.jsonl.gz'])
paragraph(140, '原v3执行存在三项方案偏离：检查点改为累计事件首次达到各十分之一及最终操作，并实时读取视图；首次运行出现历史状态落盘问题，旧稿采用关闭落盘后重跑的时延；原代码的NAV asOf实际早于提交区块60秒，旧稿未披露。v3.1保留前两项历史说明，并纠正第三项后完整重跑9次；旧结果不覆盖、不混用。当前套件顺序运行并记录实际硬件，未保证主机独占，也不选择性替换时延。源码、构建和文件清单校验未完成的运行不得进入正文。', ['spec-v3/review-amendment.v3.1.md', 'run.json#/timingScope', 'manifest.json'])
last = B['aggregate'][-1]
paragraph(141, f"在本混合负载的三个规模点上，槽数与索引随事件数增加；单位事件平均gas见表11，这不足以证明每类单项操作的成本不随历史长度变化。10,000事件目标的纯回放为{ms(last['replayMs'], 2)} ms，日志获取为{ms(last['getLogsMs'], 2)} ms。后者在本机均值上更高，但两项不等同于端到端审计时延，仍需解码、校验与业务规则分析。结果来自单基金、本地节点和未持久化历史状态的顺序负载，不能推断生产吞吐量、网络共识性能或普遍优于集中式系统。", ['B-scale-summary.json#/aggregate/2'])
paragraph(143, 'EQ5采用基于公开规则的案例验证（case-based validation），尚无独立专家评审或生产部署。该程式化开放式基金综合FATF建议10的持续尽调要求、美国Rule 22c-1(a)的下一次NAV定价及(b)/(d)的计算时点安排、UCITS第84条关于特殊情形下暂停及通知的规定，以及CSSF 24/856关于重大NAV错误及其后果处理的要求。不同法域规则用于提取运营关注点，不构成某一真实基金的完整适法性评估。案例包含管理公司、登记代理、估值代理、董事会和三名投资者；表12保留原26步。其链上结果与审计判断分别报告，数值从本次运行生成。', ['C-case-summary.json#/steps', 'spec-v3/review-amendment.v3.1.md#/Source grounding'])
paragraph(146, f"原26步全部与脚本固定的预期结果一致，其中{C['acceptedOperations']}笔被接受、{len(C['rejectedAttempts'])}次eth_call预检拒绝，产生{C['lifecycleEvents']}个生命周期事件。Q1修订为下一适用NAV匹配检查：storedAt作为本案例中计算与批准时点的代理，历史更正不作为新的常规定价批次，暂停期收到的指令使用D4 17:00的恢复估值，份额执行在17:10恢复。这些是公开的案例政策输入，无法由事件自身证明。表13同时报告原日期代理判据与修订结果；Q2–Q7保留原问题。本次答案符合补充方案，Q1不得称为原判据未改动的预注册验证。", ['C-case-summary.json', 'case/replay.json.gz#/pricingPolicy', 'spec-v3/review-amendment.v3.1.md'])
paragraph(147, '表13 案例审计查询结果（事件回放及已披露案例政策）')
cell(148, 1, 1, '是否使用请求后下一个适用常规NAV，而非仅检查估值日期晚于请求')
cell(148, 1, 2, f"{C['answers']['Q1']['matchingCasePolicy']}/{C['answers']['Q1']['total']}笔匹配；C提前受理被标记。原asOf日期代理也为{C['answers']['Q1']['originalProxyMatching']}/{C['answers']['Q1']['total']}，但两者不等价。另行链上反例中的请求跳过首个NAV，原代理通过，新审计识别为skipped-next-NAV")
overpaid = int(C['answers']['Q2']['affected'][0]['overpayment']) / 10**18
cell(148, 2, 2, f'仅A的赎回请求0；按同日更正NAV计算的差额为多付{n(overpaid)}货币单位。该算术差额不等于已确定的法律补偿责任或已执行补救')
paragraph(149, '案例支持对资格、暂停、请求和版本引用之间关系的分析，也暴露实现缺口。提前使用旧NAV和跳过下一NAV均可被冻结合约接受，事后审计必须结合明确政策；链上授权写入不能证明实际估值计算或业务批准。资格撤销后无法新申请赎回，受控退出尚未实现；补充探测仍观察到撤销资格的持有人可以向合格接收方转让，构成不符合模型双边资格条件的退出路径。更晚估值日登记后，C26的历史更正被拒绝，限制了历史错误登记。补偿、现金交割与监管通知仍在链下；本案例仅计算差额并定位受影响业务，不测试完整补救规则。模拟拒绝不进入事件流，其调用及错误另行归档。', ['C-case-summary.json', 'D-posthoc-revoked-transfer.json', 'E-next-nav-counterexample.json', 'chains/*.jsonl.gz'])
paragraph(151, '上述证据分别回答五个评价问题：所选场景下的具体实现一致性（EQ1）；相对于自建余额及许可型代币基线的约束增量（EQ2）；在声明的输入假设下独立恢复实现投影（EQ3）；本地顺序负载的描述性gas、存储和回放成本（EQ4）；所选公开规则下的业务映射及其缺口（EQ5）。案例不能证明领域完整性，混合负载平均值不能证明每类操作的历史无关性。专家评价、集中式工作流比较、受控退出、显式适用版本选择及历史更正机制仍需后续工作。', ['A-summary.json', 'B-scale-summary.json', 'C-case-summary.json'])

for index, title, sources in [(101, '3', ['spec-v3/evaluation-plan.v3.md', 'spec-v3/review-amendment.v3.1.md']),
    (105, '4', ['run.json', 'build/*.artifact.json', 'build/*-forge.json', 'A-summary.json#/productionGas', 'A-summary.json#/comparatorGas']),
    (108, '5', ['contracts/src/FundToken.sol', 'contracts/src/NAVRegistry.sol']),
    (113, '6', ['../../v2.1/m0-run-summary.local.json#/counts', 'experiments/paper85-evaluation/spec/scenarios.v1.json']),
    (118, '7', ['../../v2.1/ablation-run-summary.local.json#/counts']),
    (121, '8', ['A-summary.json#/scenarioComparison/byPredicate', 'A-scenario-rows.json']),
    (128, '9', ['A-summary.json#/recoverability']),
    (133, '10', ['A-summary.json#/productionGas/operations', 'A-summary.json#/comparatorGas']),
    (138, '11', ['B-scale-summary.json#/aggregate']),
    (145, '12', ['C-case-summary.json#/steps', 'chains/*.jsonl.gz']),
    (148, '13', ['C-case-summary.json#/answers', 'E-next-nav-counterexample.json#/pricing'])]:
    table_record(index, title, sources)
claims['tables']['11']['rowMetrics'] = [{'key': key, 'digits': digits, 'divisor': scale} for key, digits, scale in scale_rows]
claims['tables']['10']['formula'] = 'Mean over five fresh trajectories; paired workflows sum request and effect, other paired categories average the two operation means; ratio divides by B2 mean.'
claims['historicalInputs'] = {str(p.relative_to(root)): hash_file(p) for p in old_dir.glob('*.json')}
claims['historicalManifestSha256'] = hash_file(old_dir / 'sha256.txt')

# Keep the reference typography/borders, but do not inherit the manuscript's
# double-spaced body leading inside Chapter 5 tables. This prevents orphan rows.
for index in [101, 105, 108, 113, 118, 121, 128, 133, 138, 145, 148]:
    table = Table(blocks[index], doc)
    margins = table._tbl.tblPr.find(qn('w:tblCellMar'))
    if margins is None:
        margins = OxmlElement('w:tblCellMar')
        table._tbl.tblPr.append(margins)
    for edge in ['top', 'bottom', 'left', 'right']:
        element = margins.find(qn(f'w:{edge}'))
        if element is None:
            element = OxmlElement(f'w:{edge}')
            margins.append(element)
        element.set(qn('w:w'), '90')
        element.set(qn('w:type'), 'dxa')
    for row_index, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        if trpr.find(qn('w:cantSplit')) is None:
            trpr.append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = row_index < len(table.rows) - 1
                ppr = p._p.get_or_add_pPr()
                snap = ppr.find(qn('w:snapToGrid'))
                if snap is None:
                    snap = OxmlElement('w:snapToGrid')
                    ppr.append(snap)
                snap.set(qn('w:val'), '0')
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
    Paragraph(blocks[index - 1], doc).paragraph_format.keep_with_next = True
    if blocks[index + 1].tag == qn('w:p'):
        Paragraph(blocks[index + 1], doc).paragraph_format.space_before = Pt(6)

# Use native CJK family names so the bundled headless renderer displays the
# same serif/sans choices. Text, math and content outside Chapter 5 are retained.
for tree in [doc._element, doc.styles.element]:
    for fonts in tree.iter(qn('w:rFonts')):
        for attr in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            value = fonts.get(qn(f'w:{attr}'))
            if value == '宋体':
                fonts.set(qn(f'w:{attr}'), 'Songti SC' if attr == 'eastAsia' else 'Times New Roman')
            elif value == '黑体':
                fonts.set(qn(f'w:{attr}'), 'Heiti SC')

assert original_text == [''.join(e.itertext()) for e in blocks[:98] + blocks[152:]], 'Outside-Chapter-5 content changed'
# Original trailing empty paragraphs create a blank page after Chapter 7.
last_heading = blocks[153]
for element in list(doc._element.body)[list(doc._element.body).index(last_heading) + 1:]:
    if element.tag == qn('w:p') and not ''.join(element.xpath('.//w:t/text()')).strip():
        element.getparent().remove(element)
output.parent.mkdir(parents=True, exist_ok=True)
doc.save(output)
claims['outputSha256'] = hash_file(output)
claims['legalSources'] = {
    'FATF R.10': 'https://www.fatf-gafi.org/en/publications/Fatfrecommendations/Fatf-recommendations.html',
    'Rule 22c-1(a)(b)(d)': 'https://www.ecfr.gov/current/title-17/chapter-II/part-270/section-270.22c-1',
    'UCITS Art.84 / amendment EU 2024/927': 'https://eur-lex.europa.eu/eli/dir/2024/927/oj/eng',
    'CSSF 24/856 Chapter 4': 'https://www.cssf.lu/wp-content/uploads/CSSF24_856eng.pdf',
}
(output.parent / 'paper85-section5-claims.json').write_text(json.dumps(claims, ensure_ascii=False, indent=2) + '\n')
print(json.dumps({'output': str(output), 'sha256': claims['outputSha256'], 'tables': len(doc.tables), 'chapter5Tables': len(claims['tables'])}))
